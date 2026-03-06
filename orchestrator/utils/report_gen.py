import pandas as pd
import io
import logging
from datetime import datetime
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

logger = logging.getLogger(__name__)

# Define Styles
HEADER_FILL = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=12)
CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_ALIGN = Alignment(horizontal="left", vertical="top", wrap_text=True)

THIN_BORDER = Border(
    left=Side(style='thin', color='BFBFBF'),
    right=Side(style='thin', color='BFBFBF'),
    top=Side(style='thin', color='BFBFBF'),
    bottom=Side(style='thin', color='BFBFBF')
)

def get_severity_fill(severity: str) -> PatternFill:
    sev = str(severity).upper()
    if sev == "CRITICAL":
        return PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid") # Light Red
    elif sev == "HIGH":
        return PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid") # Light Orange/Yellow
    elif sev == "MEDIUM":
        return PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # Light Yellow
    elif sev == "LOW":
        return PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") # Light Green
    return PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid") # Gray Info

def apply_header_style(worksheet):
    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER

def generate_excel_report(scan_data, findings) -> bytes:
    """
    Generates a multi-sheet Excel report for the scan with professional formatting.
    Returns the Excel file as bytes.
    """
    try:
        logger.info(f"Starting Excel generation for target: {getattr(scan_data, 'target_url', 'unknown')}")
        
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            
            # --- Sheet 1: Summary ---
            summary_data = {
                "Metric": [
                    "Target URL", 
                    "Scan Date", 
                    "Scan Type", 
                    "Critical Findings", 
                    "High Findings",
                    "Total Findings"
                ],
                "Value": [
                    getattr(scan_data, 'target_url', 'Unknown'),
                    getattr(scan_data, 'created_at', datetime.now()).strftime('%Y-%m-%d %H:%M:%S'),
                    getattr(scan_data, 'scan_type', 'Unknown'),
                    getattr(scan_data, 'critical_count', 0),
                    getattr(scan_data, 'high_count', 0),
                    len(findings) if findings else 0
                ]
            }
            df_summary = pd.DataFrame(summary_data)
            df_summary.to_excel(writer, sheet_name='Summary', index=False)
            
            # Style Summary Sheet
            ws_summary = writer.sheets['Summary']
            apply_header_style(ws_summary)
            ws_summary.column_dimensions['A'].width = 25
            ws_summary.column_dimensions['B'].width = 60
            
            for row in ws_summary.iter_rows(min_row=2, max_row=ws_summary.max_row, min_col=1, max_col=2):
                for cell in row:
                    cell.alignment = LEFT_ALIGN
                    cell.border = THIN_BORDER

            # --- Sheet 2: Findings ---
            if findings:
                rows = []
                for f in findings:
                    rows.append({
                        "Severity": str(getattr(f, 'severity', 'info')).upper(),
                        "Title": getattr(f, 'title', 'Untitled'),
                        "Location": getattr(f, 'location', '-'),
                        "Description": getattr(f, 'description', ''),
                        "Tool": getattr(f, 'tool', 'Unknown')
                    })
                df_findings = pd.DataFrame(rows)
            else:
                df_findings = pd.DataFrame(columns=["Severity", "Title", "Location", "Description", "Tool"])
                
            df_findings.to_excel(writer, sheet_name='Findings', index=False)
            
            # Style Findings Sheet
            ws_findings = writer.sheets['Findings']
            apply_header_style(ws_findings)
            ws_findings.column_dimensions['A'].width = 15  # Severity
            ws_findings.column_dimensions['B'].width = 40  # Title
            ws_findings.column_dimensions['C'].width = 40  # Location
            ws_findings.column_dimensions['D'].width = 80  # Description
            ws_findings.column_dimensions['E'].width = 15  # Tool
            
            # Apply color coding and borders to data rows
            for row in ws_findings.iter_rows(min_row=2, max_row=ws_findings.max_row, min_col=1, max_col=5):
                severity_val = str(row[0].value).upper() if row[0].value else "INFO"
                fill = get_severity_fill(severity_val)
                
                for cell in row:
                    cell.alignment = LEFT_ALIGN
                    cell.border = THIN_BORDER
                # Color code only the Severity column
                row[0].fill = fill
                row[0].font = Font(bold=True)

        output.seek(0)
        return output.getvalue()

    except Exception as e:
        logger.error(f"Excel generation failed: {e}", exc_info=True)
        raise e
